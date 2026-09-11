"""
reports.py
Builds PDF (ReportLab) and DOCX (python-docx) inspection reports.

Typical flow for whoever wires the demo:
    from reports import build_pdf_report, build_docx_report
    from hashing import hash_bytes
    from storage import upload_bytes, build_key

    pdf_bytes = build_pdf_report(inspection_data)
    file_hash = hash_bytes(pdf_bytes)
    key = build_key(inspection_id, "report.pdf", prefix="reports")
    upload_bytes(pdf_bytes, key, content_type="application/pdf")
    # then save Report(inspection_id=..., report_type="pdf",
    #                   file_path=key, evidence_hash=file_hash)
"""

import io
import os
from datetime import datetime
from typing import Any, Dict, List, Optional

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

# ---------------------------------------------------------------------------
# Unicode font registration.
#
# ReportLab's built-in fonts (Helvetica etc.) don't include the ₹ (Rupee)
# glyph or other non-Latin-1 characters — they render as a black box (■).
# DejaVu Sans does include ₹, so we bundle it and register it here.
# Falls back to Helvetica (with a printed warning) if the font file is
# missing, so the report still generates, just without the ₹ glyph.
# ---------------------------------------------------------------------------
_FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "fonts")
_REGULAR_PATH = os.path.join(_FONT_DIR, "DejaVuSans.ttf")
_BOLD_PATH = os.path.join(_FONT_DIR, "DejaVuSans-Bold.ttf")

try:
    pdfmetrics.registerFont(TTFont("DejaVuSans", _REGULAR_PATH))
    pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", _BOLD_PATH))
    BODY_FONT = "DejaVuSans"
    BOLD_FONT = "DejaVuSans-Bold"
except Exception as e:  # font files missing or unreadable
    print(f"WARNING: could not load DejaVuSans font ({e}); falling back to "
          f"Helvetica. Currency symbols like ₹ may render incorrectly.")
    BODY_FONT = "Helvetica"
    BOLD_FONT = "Helvetica-Bold"


# ---------------------------------------------------------------------------
# Expected shape of `inspection_data` passed into both builders:
#
# {
#   "inspection_id": 42,
#   "source": "amazon.in/listing/xyz",
#   "input_mode": "bulk",
#   "status": "needs_review",
#   "created_at": datetime(...),
#   "extracted_fields": [
#       {"field_name": "mrp", "field_value": "₹199", "confidence": 0.94},
#       ...
#   ],
#   "violations": [
#       {"rule_code": "LM-PCR-2011-R6", "description": "...",
#        "severity": "medium", "status": "needs_review"},
#       ...
#   ],
#   "evidence_hashes": ["a1b2c3...", "d4e5f6..."],
# }
# ---------------------------------------------------------------------------


def build_pdf_report(inspection_data: Dict[str, Any]) -> bytes:
    """Builds a PDF inspection report in memory and returns its bytes."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    # Point the styles we use at the Unicode font so ₹ and other
    # non-Latin-1 characters render instead of showing as ■.
    styles["Title"].fontName = BOLD_FONT
    styles["Heading2"].fontName = BOLD_FONT
    styles["Normal"].fontName = BODY_FONT
    if "Code" in styles:
        styles["Code"].fontName = BODY_FONT
    story: List[Any] = []

    story.append(Paragraph("MetroScan AI — Inspection Report", styles["Title"]))
    story.append(Spacer(1, 12))

    meta_lines = [
        f"Inspection ID: {inspection_data.get('inspection_id')}",
        f"Source: {inspection_data.get('source', 'N/A')}",
        f"Input mode: {inspection_data.get('input_mode', 'N/A')}",
        f"Status: {inspection_data.get('status', 'N/A')}",
        f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}",
    ]
    for line in meta_lines:
        story.append(Paragraph(line, styles["Normal"]))
    story.append(Spacer(1, 16))

    # Small helper so long text wraps inside table cells instead of
    # overflowing into the next column.
    cell_style = styles["Normal"]

    def cell(text: str) -> Paragraph:
        return Paragraph(str(text), cell_style)

    # --- Extracted fields table ---
    story.append(Paragraph("Extracted Fields", styles["Heading2"]))
    fields = inspection_data.get("extracted_fields", [])
    if fields:
        table_data = [["Field", "Value", "Confidence"]]
        for f in fields:
            confidence = f.get("confidence")
            conf_str = f"{confidence:.0%}" if confidence is not None else "N/A"
            table_data.append([cell(f.get("field_name", "")), cell(f.get("field_value", "")), cell(conf_str)])
        table = Table(table_data, colWidths=[5 * cm, 7 * cm, 3 * cm])
        table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), BODY_FONT),
                    ("FONTNAME", (0, 0), (-1, 0), BOLD_FONT),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E5F")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F0F0")]),
                ]
            )
        )
        story.append(table)
    else:
        story.append(Paragraph("No fields extracted.", styles["Normal"]))
    story.append(Spacer(1, 16))

    # --- Rule violations table ---
    story.append(Paragraph("Rule Violations", styles["Heading2"]))
    violations = inspection_data.get("violations", [])
    if violations:
        v_data = [["Rule Code", "Description", "Severity", "Status"]]
        for v in violations:
            v_data.append(
                [
                    cell(v.get("rule_code", "")),
                    cell(v.get("description", "")),
                    cell(v.get("severity", "")),
                    cell(v.get("status", "")),
                ]
            )
        v_table = Table(v_data, colWidths=[3.5 * cm, 6.5 * cm, 2.5 * cm, 2.5 * cm])
        v_table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), BODY_FONT),
                    ("FONTNAME", (0, 0), (-1, 0), BOLD_FONT),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#B3401D")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F0F0")]),
                ]
            )
        )
        story.append(v_table)
    else:
        story.append(Paragraph("No violations detected.", styles["Normal"]))
    story.append(Spacer(1, 16))

    # --- Evidence integrity ---
    story.append(Paragraph("Evidence Integrity (SHA-256)", styles["Heading2"]))
    hashes = inspection_data.get("evidence_hashes", [])
    if hashes:
        for h in hashes:
            story.append(Paragraph(h, styles["Code"] if "Code" in styles else styles["Normal"]))
    else:
        story.append(Paragraph("No evidence hashes recorded.", styles["Normal"]))

    doc.build(story)
    return buffer.getvalue()


def build_docx_report(inspection_data: Dict[str, Any]) -> bytes:
    """Builds a DOCX inspection report in memory and returns its bytes."""
    document = Document()

    document.add_heading("MetroScan AI — Inspection Report", level=0)

    meta = document.add_paragraph()
    meta.add_run(f"Inspection ID: {inspection_data.get('inspection_id')}\n").bold = False
    meta.add_run(f"Source: {inspection_data.get('source', 'N/A')}\n")
    meta.add_run(f"Input mode: {inspection_data.get('input_mode', 'N/A')}\n")
    meta.add_run(f"Status: {inspection_data.get('status', 'N/A')}\n")
    meta.add_run(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")

    # --- Extracted fields ---
    document.add_heading("Extracted Fields", level=1)
    fields = inspection_data.get("extracted_fields", [])
    if fields:
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Field", "Value", "Confidence"
        for f in fields:
            row = table.add_row().cells
            row[0].text = str(f.get("field_name", ""))
            row[1].text = str(f.get("field_value", ""))
            confidence = f.get("confidence")
            row[2].text = f"{confidence:.0%}" if confidence is not None else "N/A"
    else:
        document.add_paragraph("No fields extracted.")

    # --- Rule violations ---
    document.add_heading("Rule Violations", level=1)
    violations = inspection_data.get("violations", [])
    if violations:
        table = document.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 2"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = (
            "Rule Code",
            "Description",
            "Severity",
            "Status",
        )
        for v in violations:
            row = table.add_row().cells
            row[0].text = str(v.get("rule_code", ""))
            row[1].text = str(v.get("description", ""))
            row[2].text = str(v.get("severity", ""))
            row[3].text = str(v.get("status", ""))
    else:
        document.add_paragraph("No violations detected.")

    # --- Evidence integrity ---
    document.add_heading("Evidence Integrity (SHA-256)", level=1)
    hashes = inspection_data.get("evidence_hashes", [])
    if hashes:
        for h in hashes:
            p = document.add_paragraph(h)
            p.runs[0].font.size = Pt(9)
    else:
        document.add_paragraph("No evidence hashes recorded.")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()